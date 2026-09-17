"""生成《明胶分子量复配计算与AI辅助决策Demo项目报价方案》Word 文档。"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_FILE = Path(__file__).with_name(
    "明胶分子量复配计算与AI辅助决策Demo项目报价方案.docx"
)
CHINESE_FONT = "SimSong"


def set_cell_shading(cell, fill: str) -> None:
    """设置表格单元格背景色。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text: str, *, bold: bool = False, color=None) -> None:
    """填充表格文字并统一字体。"""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = CHINESE_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    """让跨页表格重复显示表头。"""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    for index, title in enumerate(headers):
        set_cell_text(header.cells[index], title, bold=True, color=(255, 255, 255))
        set_cell_shading(header.cells[index], "1F4E78")
        if widths:
            header.cells[index].width = Cm(widths[index])
    set_repeat_table_header(header)

    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            set_cell_text(row.cells[index], str(value))
            if widths:
                row.cells[index].width = Cm(widths[index])
    document.add_paragraph()
    return table


def add_heading(document, text: str, level: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = f"Heading {level}"
    run = paragraph.add_run(text)
    run.font.name = CHINESE_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    run.font.color.rgb = RGBColor(31, 78, 120)


def add_body(document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.35
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        prefix.bold = True
        prefix.font.name = CHINESE_FONT
        prefix._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
        remainder = paragraph.add_run(text[len(bold_prefix):])
        remainder.font.name = CHINESE_FONT
        remainder._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    else:
        run = paragraph.add_run(text)
        run.font.name = CHINESE_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    for run in paragraph.runs:
        run.font.size = Pt(10.5)
    return paragraph


def add_bullet(document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.name = CHINESE_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    run.font.size = Pt(10.5)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    normal_style = document.styles["Normal"]
    normal_style.font.name = CHINESE_FONT
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    normal_style.font.size = Pt(10.5)

    for level in (1, 2, 3):
        style = document.styles[f"Heading {level}"]
        style.font.name = CHINESE_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
        style.font.color.rgb = RGBColor(31, 78, 120)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("明胶分子量复配计算与 AI 辅助决策 Demo 系统｜项目报价方案")
    run.font.name = CHINESE_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(127, 127, 127)


def main() -> None:
    document = Document()
    configure_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(72)
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run("明胶分子量复配计算与\nAI 辅助决策 Demo 系统")
    run.bold = True
    run.font.name = CHINESE_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor(31, 78, 120)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(54)
    run = subtitle.add_run("项目报价与交付方案")
    run.font.name = CHINESE_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(89, 89, 89)

    info = [
        "项目周期：约 2 个月（8 周）",
        "报价口径：未税，税费按双方最终合同约定另行核算",
        "版本日期：2026 年 8 月",
    ]
    for line in info:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.name = CHINESE_FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(89, 89, 89)

    document.add_section(WD_SECTION.NEW_PAGE)

    add_heading(document, "一、项目目标", 1)
    add_body(
        document,
        "将明胶生产中已有的分子量、黏度及复配经验公式代码化，形成可演示、可调整、可复核的网页工具。系统用于辅助工艺人员完成 GPC 分子量计算、黏度法估算、胶液复配计算、目标偏差判断和规则式工艺建议，降低人工计算与经验判断的依赖。",
    )
    add_body(
        document,
        "本期系统定位为工艺计算与辅助决策 Demo，不包含机器学习预测模型开发；系统输出不替代实验检测、生产投料决策或最终质量判定。",
    )

    add_heading(document, "二、总体报价与付款方式", 1)
    add_table(
        document,
        ["付款节点", "比例", "金额（未税）", "触发条件"],
        [
            ["项目启动款", "30%", "108,000 元", "合同签订后 5 个工作日内支付"],
            ["Demo 交付款", "20%", "72,000 元", "Demo 开发完成、演示交付并阶段确认后支付"],
            ["正式验收款", "50%", "180,000 元", "试运行优化完成、正式验收通过后支付"],
            ["开发费合计", "100%", "360,000 元", "项目开发总价"],
            ["上线后运维服务费", "—", "8,000 元/月", "正式验收后次月起按月结算"],
        ],
        widths=[3.1, 1.4, 3.0, 8.0],
    )
    add_body(document, "建议对外报价为 36 万元未税。", bold_prefix="建议对外报价为 36 万元未税。")
    add_body(
        document,
        "如客户预算压力较大，可在不改变核心算法范围的前提下协商至 30 万元；如需部署、账号权限、数据导入、报表导出、接口集成等扩展能力，建议按 40 万元或以上另行报价。",
    )

    add_heading(document, "三、阶段一：Demo 开发与演示交付", 1)
    add_body(document, "周期：第 1–4 周；对应费用：180,000 元（含 108,000 元启动款及 72,000 元 Demo 交付款）。")
    add_heading(document, "1. 建设内容", 2)
    for item in [
        "GPC 分子量计算：Mn、Mw、Mz、PDI，以及小分子肽占比计算。",
        "黏度法 Mw 估算，支持 K、a 参数调整。",
        "多批胶液复配 Mw 计算，支持按固含量折算有效干基重量。",
        "目标 Mw、偏差、偏差率、合格范围及偏高/偏低状态判断。",
        "GPC、黏度法、复配三种 Web 可视化计算模式。",
        "批次名称、胶液类型、重量、固含量、Mw 等参数动态维护与实时计算。",
        "标准计算公式展示、结果卡片、规则引擎式 AI 提问与调配建议。",
        "基础单元测试、示例数据、启动脚本和操作说明。",
    ]:
        add_bullet(document, item)
    add_heading(document, "2. 阶段交付物", 2)
    add_table(
        document,
        ["类别", "交付内容"],
        [
            ["可运行系统", "明胶分子量复配计算与 AI 辅助决策 Demo 页面"],
            ["源代码", "app.py、calculator.py、tests.py 等完整源代码"],
            ["算法与测试材料", "公式说明、测试用例、测试结果与示例数据"],
            ["使用材料", "README、启动脚本、基础操作说明"],
            ["阶段确认材料", "Demo 演示记录或阶段确认单"],
        ],
        widths=[3.2, 12.3],
    )
    add_heading(document, "3. 阶段确认建议", 2)
    for item in [
        "系统可在约定电脑或测试环境中正常启动。",
        "指定测试案例的计算结果与人工复核结果一致。",
        "可录入、修改并动态计算多批胶液参数。",
        "可展示目标 Mw、当前 Mw、偏差、偏差率和规则式建议。",
        "客户完成 Demo 演示确认。",
    ]:
        add_bullet(document, item)

    add_heading(document, "四、阶段二：试运行优化、正式交付与验收", 1)
    add_body(document, "周期：第 5–8 周；对应费用：180,000 元（正式验收款）。")
    add_heading(document, "1. 建设内容", 2)
    for item in [
        "协助客户使用真实或脱敏工艺样本进行试运行，并复核计算结果和参数填写方式。",
        "收集工艺、质控及管理人员的使用反馈，形成试运行问题清单。",
        "在已确认范围内修正公式实现或显示问题，优化字段、默认值、提示文案和页面布局。",
        "优化规则式 AI 提问与复配方向建议，补充双方确认的常规测试数据和操作说明。",
        "提供正式版本部署包或源代码包，完成一次系统使用培训/交接。",
        "提供交付清单、验收测试清单、版本说明，并配合正式验收。",
    ]:
        add_bullet(document, item)
    add_heading(document, "2. 阶段交付物", 2)
    add_table(
        document,
        ["类别", "交付内容"],
        [
            ["正式版本", "根据试运行反馈优化后的系统版本"],
            ["配置材料", "参数配置说明、默认值说明及测试数据模板"],
            ["验收材料", "验收测试清单、测试结果、问题闭环记录"],
            ["使用与培训材料", "更新后的用户操作手册/培训材料及培训记录"],
            ["正式交付材料", "源代码或部署包、版本说明、交付清单、验收单"],
        ],
        widths=[3.2, 12.3],
    )
    add_heading(document, "3. 正式验收建议", 2)
    for item in [
        "阶段一功能在正式版本中稳定可用。",
        "试运行期间双方确认的问题已完成整改或形成书面处理结论。",
        "关键测试数据的计算结果符合双方已确认的公式。",
        "客户指定人员可独立完成基础操作。",
        "双方签署正式验收单。",
    ]:
        add_bullet(document, item)

    add_heading(document, "五、项目计划", 1)
    add_table(
        document,
        ["周期", "主要工作", "关键输出"],
        [
            ["第 1 周", "工艺公式、字段、测试样本确认", "需求确认清单、计算口径确认"],
            ["第 2–3 周", "算法开发、页面开发、基础测试", "可运行 Demo 初版"],
            ["第 4 周", "Demo 联调、演示、阶段交付", "Demo 交付与阶段确认"],
            ["第 5–6 周", "客户试运行、收集反馈", "试运行问题清单"],
            ["第 7 周", "范围内优化、回归测试", "正式版本候选包"],
            ["第 8 周", "培训、正式交付、验收", "正式验收与项目交付"],
        ],
        widths=[2.0, 6.7, 6.8],
    )
    add_body(document, "项目周期以客户及时提供工艺公式、测试数据、试运行反馈及验收安排为前提。因客户侧数据、接口、环境或确认延迟造成的等待时间，相应顺延。")

    add_heading(document, "六、上线后运维服务", 1)
    add_body(document, "服务费用：8,000 元/月（未税）；服务起算：正式验收后的次月起；建议首签服务周期：12 个月。")
    add_heading(document, "1. 运维服务包含", 2)
    for item in [
        "系统运行咨询与远程支持。",
        "已交付功能的缺陷修复。",
        "页面文案、默认参数、规则提示等轻量级配置调整。",
        "月度一次运行情况沟通或问题汇总。",
        "常规依赖升级风险评估与必要的兼容性处理。",
    ]:
        add_bullet(document, item)
    add_heading(document, "2. 运维服务不包含", 2)
    for item in [
        "新增机器学习/预测模型开发或真实大模型接入。",
        "ERP、MES、LIMS、GPC 仪器或其他第三方系统接口开发。",
        "用户、角色、权限、单点登录等企业级账号体系建设。",
        "数据库、历史数据治理、报表中心或移动端开发。",
        "云服务器、域名、SSL 证书、数据库及第三方软件许可费用。",
        "超出本期范围的新页面、新算法、新报表或重大交互改造。",
    ]:
        add_bullet(document, item)
    add_body(document, "上述新增需求可按人天、按模块或另行项目报价。")

    add_heading(document, "七、双方配合事项与范围边界", 1)
    add_heading(document, "1. 客户方需提供", 2)
    for item in [
        "已确认的工艺公式、参数范围和单位口径。",
        "典型批次样本及人工复核结果。",
        "目标 Mw、允许偏差及质量判断规则。",
        "试运行人员、验收人员和反馈时限。",
        "如需部署到客户环境，提供相应电脑/服务器及网络条件。",
    ]:
        add_bullet(document, item)
    add_heading(document, "2. 范围边界与变更原则", 2)
    for item in [
        "本期仅实现明确的公式计算、参数调整、可视化展示和规则式辅助建议。",
        "系统输出为工艺辅助结果，最终生产投料和质量判定仍由客户工艺及质控人员负责。",
        "工艺公式调整、新增质量指标、新增数据源、第三方接口或部署架构变化导致的新增工作，双方应书面确认变更范围、工期与费用。",
        "建议将客户确认的公式版本、测试数据和验收标准作为合同附件，以避免后续计算口径争议。",
    ]:
        add_bullet(document, item)

    add_heading(document, "八、报价结论", 1)
    add_body(document, "本项目建议采用“36 万元开发费 + 8,000 元/月运维费”的合作方式。")
    add_table(
        document,
        ["节点", "付款金额", "说明"],
        [
            ["合同签订、项目启动", "108,000 元", "启动款，覆盖需求确认及前期开发投入"],
            ["Demo 开发完成并阶段确认", "72,000 元", "Demo 交付款"],
            ["试运行优化完成、正式验收", "180,000 元", "正式验收款"],
            ["正式验收后", "8,000 元/月", "持续运维与技术支持"],
        ],
        widths=[4.2, 3.1, 8.2],
    )
    add_body(document, "该方案兼顾项目启动投入、客户先看到可运行成果再持续付款的诉求，并为试运行优化和正式验收预留了清晰的工作与结算空间。")

    document.save(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()
