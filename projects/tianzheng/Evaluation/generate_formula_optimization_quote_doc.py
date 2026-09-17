"""生成明胶半成品配方优化系统报价方案 Word 文档。"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_FILE = Path(__file__).with_name("明胶半成品配方优化系统项目报价方案.docx")
FONT = "STHeiti"
BLUE = "1F4E78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
USABLE_WIDTH_DXA = 9360


def set_run_font(run, size=10.5, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)
    lang = run._element.rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.rPr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    lang.set(qn("w:bidi"), "zh-CN")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    cell_mar = tc_pr.first_child_found_in("w:tcMar")
    if cell_mar is None:
        cell_mar = OxmlElement("w:tcMar")
        tc_pr.append(cell_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = cell_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margin(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(str(text))
    set_run_font(run, size=9.5, bold=bold, color=color)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(USABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    header = table.rows[0]
    for cell, text in zip(header.cells, headers):
        set_cell_text(cell, text, bold=True, color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, BLUE)
    repeat_header(header)
    for row_values in rows:
        row = table.add_row()
        for cell, text in zip(row.cells, row_values):
            set_cell_text(cell, text)
    set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    return table


def add_text(doc, text, *, size=10.5, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 12}[level], bold=True, color=(46, 116, 181) if level < 3 else (31, 77, 120))
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.95)
    p.paragraph_format.first_line_indent = Cm(-0.47)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    for level, size, before, after, color in (
        (1, 16, 18, 10, (46, 116, 181)),
        (2, 13, 12, 6, (46, 116, 181)),
        (3, 12, 8, 4, (31, 77, 120)),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("明胶半成品配方优化系统｜项目报价与交付方案")
    set_run_font(run, size=8.5, color=(127, 127, 127))


def main():
    doc = Document()
    configure_document(doc)

    # proposal_centerpiece opening block
    add_text(doc, "项目报价与交付方案", size=12, bold=True, color=(89, 89, 89), align=WD_ALIGN_PARAGRAPH.CENTER, before=38, after=10)
    add_text(doc, "明胶半成品配方优化系统", size=25, bold=True, color=(31, 78, 120), align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_text(doc, "从加权平均计算升级为多方案配方优化与 Excel 批次数据维护", size=12.5, color=(89, 89, 89), align=WD_ALIGN_PARAGRAPH.CENTER, after=28)
    add_table(
        doc,
        ["项目周期", "开发报价（未税）", "付款方式", "运维服务"],
        [["约 2 个月（8 周）", "360,000 元", "30% / 20% / 50%", "8,000 元/月"]],
        [2200, 2400, 2400, 2360],
    )
    add_text(doc, "报价说明：税费按双方最终合同约定另行核算。", size=9.5, color=(89, 89, 89), align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    add_heading(doc, "一、项目背景与建设目标", 1)
    add_text(doc, "客户当前针对明胶半成品复配主要采用加权平均方式进行配方计算，难以在多个批次、库存、目标指标及工艺约束同时存在时快速得到可执行的优选方案。")
    add_text(doc, "本项目拟建设“明胶半成品配方优化系统”：以现有配方计算逻辑为基础，接入 Excel 批次数据，通过可配置的目标与约束条件，自动生成多套候选配方并排序推荐，帮助工艺人员进行配方选择、人工微调与结果复核。")
    add_text(doc, "本期定位：", bold=True, after=2)
    for item in [
        "从单一加权平均计算，升级为面向目标指标和约束条件的配方优化。",
        "一次输出多套候选最优方案（建议默认 3–5 套），供工艺人员比较与选择。",
        "支持通过 Excel 模板更新半成品批次、库存与指标数据，无需每次修改程序。",
        "保留人工调整能力，便于现场结合实际情况进行二次决策。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "二、总体报价与付款方式", 1)
    add_table(
        doc,
        ["付款节点", "比例", "金额（未税）", "付款条件"],
        [
            ["项目启动款", "30%", "108,000 元", "合同签订后 5 个工作日内支付"],
            ["Demo 交付款", "20%", "72,000 元", "Demo 完成、演示交付并经阶段确认后支付"],
            ["正式验收款", "50%", "180,000 元", "试运行优化完成、正式验收通过后支付"],
            ["开发费合计", "100%", "360,000 元", "项目开发总价"],
            ["上线后运维服务", "—", "8,000 元/月", "正式验收后次月起按月结算"],
        ],
        [2100, 1000, 2200, 4060],
    )
    add_text(doc, "报价说明：本报价已包含多方案配方优化、Excel 数据导入与校验、优化结果排序及试运行参数校准等核心功能；超出本期范围的系统集成、数据库建设及新增算法功能另行评估。", size=10, color=(89, 89, 89))

    add_heading(doc, "三、阶段一：配方优化 Demo 开发与交付", 1)
    add_text(doc, "周期：第 1–4 周；阶段费用：180,000 元（含 108,000 元启动款及 72,000 元 Demo 交付款）。")
    add_heading(doc, "1. 建设内容", 2)
    for item in [
        "需求与计算口径确认：梳理半成品字段、目标指标、可用批次、库存限制、配方边界及输出规则。",
        "Excel 数据模板设计：定义批次名称、半成品类型、可用量、固含量、关键质量指标、成本（如适用）等字段；支持模板上传、数据校验与更新。",
        "配方优化引擎：以现有加权平均计算为基线，结合目标值、允许偏差、批次可用量、固含量及双方确认的约束条件，计算候选配方。",
        "多方案输出：默认生成 3–5 套满足约束的候选优选方案，展示配方用量、预计指标、目标偏差、原料占比及推荐排序。",
        "可视化 Demo 页面：支持 Excel 数据导入、目标参数设置、约束条件输入、候选方案对比、方案详情和人工调整后的实时复算。",
        "基础测试：使用双方确认的典型样本验证数据导入、加权计算、约束校验与候选方案输出。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2. 阶段一交付物", 2)
    add_table(
        doc,
        ["类别", "交付内容"],
        [
            ["可运行 Demo", "明胶半成品配方优化 Web Demo 系统"],
            ["优化能力", "多目标/多约束候选方案计算与 3–5 套方案排序展示"],
            ["Excel 模板", "半成品批次数据导入模板、字段说明及校验规则"],
            ["算法材料", "配方计算口径、优化约束说明、方案排序规则说明"],
            ["测试材料", "典型测试数据、测试结果及阶段演示记录"],
            ["使用材料", "启动说明、基础操作说明及 Demo 交付清单"],
        ],
        [2100, 7260],
    )
    add_heading(doc, "3. 阶段确认建议", 2)
    for item in [
        "可按规定格式导入 Excel 半成品批次数据，并对缺失或异常数据进行提示。",
        "可配置目标指标与基础约束条件，并完成配方计算。",
        "针对同一批次数据与目标，系统可输出不少于 3 套候选方案（在满足数据与约束条件可行的前提下）。",
        "每套方案可展示组分、投料量、预计结果、偏差及推荐理由。",
        "客户完成 Demo 演示确认。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "四、阶段二：试运行优化、正式交付与验收", 1)
    add_text(doc, "周期：第 5–8 周；阶段费用：180,000 元（正式验收款）。")
    add_heading(doc, "1. 建设内容", 2)
    for item in [
        "试运行支持：协助客户使用真实或脱敏半成品批次数据进行试运行，核对 Excel 数据口径、方案可行性与现场使用方式。",
        "范围内优化：根据双方确认的试运行反馈，优化约束条件、候选方案排序、字段展示、Excel 校验提示和人工调整交互。",
        "参数配置：固化双方确认的默认目标、允许偏差、候选方案数量及基础业务规则。",
        "正式交付：提供优化后的正式版本、Excel 模板、配置说明、操作手册、测试记录及交付清单。",
        "培训与验收：完成一次系统使用培训/交接，配合客户完成正式验收。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2. 阶段二交付物", 2)
    add_table(
        doc,
        ["类别", "交付内容"],
        [
            ["正式系统", "根据试运行反馈优化后的正式版本"],
            ["数据材料", "正式 Excel 数据模板、字段字典、导入与更新操作说明"],
            ["规则材料", "目标指标、约束条件、候选方案排序等配置说明"],
            ["验收材料", "验收测试清单、测试结果、问题闭环记录"],
            ["交付材料", "源代码或部署包、版本说明、用户操作手册、培训记录、验收单"],
        ],
        [2100, 7260],
    )
    add_heading(doc, "3. 正式验收建议", 2)
    for item in [
        "正式版本可稳定完成 Excel 数据导入、更新与校验。",
        "系统可按双方确认的目标及约束条件输出候选配方方案。",
        "候选方案的预计指标、原料使用量和偏差结果可复核。",
        "试运行中双方确认的问题已完成整改或形成书面处理结论。",
        "客户指定人员可独立完成批次数据更新、方案生成和基础操作。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "五、项目计划", 1)
    add_table(
        doc,
        ["周期", "主要工作", "关键输出"],
        [
            ["第 1 周", "确认目标指标、约束条件、Excel 字段与测试样本", "需求确认清单、数据字典、优化口径"],
            ["第 2–3 周", "Excel 导入、计算与优化引擎、页面开发", "可运行 Demo 初版"],
            ["第 4 周", "Demo 联调、测试、演示交付", "阶段确认与 Demo 交付款"],
            ["第 5–6 周", "客户试运行、收集反馈、优化规则复核", "试运行问题清单"],
            ["第 7 周", "范围内优化、回归测试、交付材料整理", "正式版本候选包"],
            ["第 8 周", "培训、正式交付与验收", "正式验收与项目交付"],
        ],
        [1200, 4700, 3460],
    )
    add_text(doc, "项目周期以客户及时提供 Excel 样本、目标与约束规则、试运行反馈及验收安排为前提。因客户侧数据、确认或环境准备延迟造成的等待时间，相应顺延。", size=10, color=(89, 89, 89))

    add_heading(doc, "六、上线后运维服务", 1)
    add_text(doc, "服务费用：8,000 元/月（未税）；服务起算：正式验收后的次月起；建议首签服务周期：12 个月。")
    add_heading(doc, "1. 运维服务包含", 2)
    for item in [
        "已交付系统的运行咨询、远程支持与缺陷修复。",
        "Excel 导入模板字段说明、数据校验提示、默认参数和展示文案等轻量级调整。",
        "已确认优化规则内的阈值、默认目标、候选方案数量等配置调整。",
        "月度一次运行情况沟通或问题汇总。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2. 运维服务不包含", 2)
    for item in [
        "新增质量指标、重构配方优化算法或新增重大约束条件。",
        "ERP、MES、LIMS、仪器设备或其他第三方系统接口开发。",
        "Excel 以外的数据库建设、历史数据治理、企业级权限体系、移动端或报表中心开发。",
        "云服务器、域名、SSL 证书、数据库及第三方软件许可费用。",
        "机器学习预测模型、真实大模型接入及自动化生产控制功能。",
    ]:
        add_bullet(doc, item)
    add_text(doc, "上述新增需求可按人天、按模块或另行项目报价。", size=10, color=(89, 89, 89))

    add_heading(doc, "七、范围边界与双方配合", 1)
    add_heading(doc, "1. 客户方需提供", 2)
    for item in [
        "可用于试运行的半成品批次 Excel 样本及字段含义说明。",
        "目标指标、允许偏差、库存/可用量限制及其他业务约束条件。",
        "候选方案的评价偏好或排序规则，例如目标贴合度、原料利用率、成本优先级等。",
        "典型人工配方及人工复核结果，用于验证系统输出合理性。",
        "试运行人员、验收人员和反馈时限。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2. 范围边界与变更原则", 2)
    for item in [
        "本期优化结果为辅助决策结果，最终配方确认与生产投料仍由客户工艺及质控人员负责。",
        "“最优方案”以双方确认的目标、约束和排序规则为准；当不存在满足全部约束的可行方案时，系统应给出原因提示或接近目标的候选方案。",
        "工艺规则变化、新增关键指标、第三方接口、数据库或部署架构变化导致的新增工作，双方应书面确认范围、工期与费用。",
        "建议将 Excel 模板、数据字典、目标指标、约束条件、方案排序规则及验收样本作为合同附件。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "八、报价结论", 1)
    add_text(doc, "本项目建议采用“36 万元开发费 + 8,000 元/月运维费”的合作方式。")
    add_table(
        doc,
        ["节点", "付款金额", "说明"],
        [
            ["合同签订、项目启动", "108,000 元", "启动款，覆盖需求确认、数据模板和前期开发投入"],
            ["Demo 完成并阶段确认", "72,000 元", "Excel 导入与多方案配方优化 Demo 交付款"],
            ["试运行优化完成、正式验收", "180,000 元", "正式验收款"],
            ["正式验收后", "8,000 元/月", "持续运维与技术支持"],
        ],
        [3000, 2200, 4160],
    )
    add_text(doc, "该方案以“Excel 批次数据更新 + 多套候选最优配方输出 + 试运行校准”为核心交付，适用于先完成数字化辅助决策，再逐步扩展预测模型、系统集成及生产优化能力的实施路径。")

    doc.save(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()
