"""
core/customer_exporter.py
=========================
客户关键词识别结果导出。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.config import settings


_HEADER_FILL = PatternFill("solid", fgColor="E7F0EA")
_WARN_FILL = PatternFill("solid", fgColor="FFF4D8")
_ERROR_FILL = PatternFill("solid", fgColor="FFE1E1")
_OK_FILL = PatternFill("solid", fgColor="E7F7EF")
_BOLD = Font(bold=True)


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:g}"
    return str(value)


def _status_fill(status: str):
    if status == "ok":
        return _OK_FILL
    if status in ("missing", "error"):
        return _ERROR_FILL
    return _WARN_FILL


def write_customer_result_workbook(
    *,
    file_id: str,
    fields: List[Dict[str, Any]],
    line_items: List[Dict[str, Any]] | None = None,
    business_result: Dict[str, Any],
    config: Dict[str, Any],
    source_filename: str,
) -> Path:
    """生成客户关键词结果 Excel，返回文件路径。"""
    settings.CUSTOMER_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = settings.CUSTOMER_EXPORT_DIR / f"{file_id}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "校对结果"

    ws["A1"] = "客户关键词识别结果"
    ws["A1"].font = Font(bold=True, size=15)
    ws["A2"] = "来源文件"
    ws["B2"] = source_filename
    ws["A3"] = "关键词配置"
    ws["B3"] = config.get("label", config.get("id", ""))
    ws["A4"] = "业务处理"
    ws["B4"] = business_result.get("message", "")
    ws["A5"] = "处理类型"
    ws["B5"] = business_result.get("result_type", "")
    ws["A6"] = "是否需要人工确认"
    ws["B6"] = "是" if business_result.get("requires_manual_confirmation") else "否"

    for row in range(2, 7):
        ws[f"A{row}"].font = _BOLD

    start_row = 8
    headers = [
        "字段ID",
        "字段名称",
        "原始识别值",
        "标准化值",
        "校验状态",
        "错误/提示",
        "是否必填",
        "业务主键",
        "来源",
        "输出单元格",
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _BOLD
        cell.alignment = Alignment(horizontal="center")

    for row_idx, field in enumerate(fields, start=start_row + 1):
        values = [
            field.get("id", ""),
            field.get("label", ""),
            _stringify(field.get("raw_value", "")),
            _stringify(field.get("standard_value", "")),
            field.get("validation_status", ""),
            "; ".join(field.get("errors", []) or []) or field.get("note", ""),
            "是" if field.get("required") else "否",
            "是" if field.get("business_key") else "否",
            field.get("source", ""),
            field.get("output_cell", ""),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col == 5:
                cell.fill = _status_fill(str(value))

    detail = wb.create_sheet("业务处理")
    detail_headers = ["项目", "值"]
    for col, header in enumerate(detail_headers, start=1):
        cell = detail.cell(row=1, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _BOLD

    business_rows = [
        ("主键字段", business_result.get("business_key_id", "")),
        ("主键值", business_result.get("business_key_value", "")),
        ("结果类型", business_result.get("result_type", "")),
        ("动作", business_result.get("action", "")),
        ("旧值", _stringify(business_result.get("old_value", ""))),
        ("新值", _stringify(business_result.get("new_value", ""))),
        ("匹配 MOQ", _stringify(business_result.get("matched_moq", ""))),
        ("提示", business_result.get("message", "")),
    ]
    for row, (key, value) in enumerate(business_rows, start=2):
        detail.cell(row=row, column=1, value=key).font = _BOLD
        detail.cell(row=row, column=2, value=value)

    item_sheet = None
    if line_items:
        item_sheet = wb.create_sheet("明细行")
        max_prices = max(
            1,
            max(
                len(item.get("prices", []) or [])
                for item in line_items
            ),
        )
        item_headers = [
            "行号",
            "SKU",
            "客户物料代码",
            "品名",
            "规格型号/主键",
            "材质及特殊特性",
            "硬度",
            "颜色",
            "交货日期",
            "表格来源",
        ]
        for idx in range(1, max_prices + 1):
            item_headers.extend(
                [
                    f"价格{idx}",
                    f"MOQ{idx}",
                    f"数量{idx}",
                    f"价税合计{idx}",
                    f"价格来源{idx}",
                ]
            )
        item_headers.extend(
            [
            "校验状态",
            "错误/提示",
            "来源",
            ]
        )
        for col, header in enumerate(item_headers, start=1):
            cell = item_sheet.cell(row=1, column=col, value=header)
            cell.fill = _HEADER_FILL
            cell.font = _BOLD
            cell.alignment = Alignment(horizontal="center")

        for row_idx, item in enumerate(line_items, start=2):
            values = [
                item.get("line_no", ""),
                item.get("sku", ""),
                item.get("item_no", ""),
                item.get("product_name", ""),
                item.get("part_no", ""),
                item.get("material_special", ""),
                item.get("hardness", ""),
                item.get("color", ""),
                item.get("lead_time", ""),
                item.get("source_table", ""),
            ]
            prices = item.get("prices", []) or []
            if not prices:
                prices = [
                    {
                        "tax_included_price": item.get("tax_included_price", ""),
                        "moq": item.get("moq", ""),
                        "quantity": item.get("quantity", ""),
                        "amount": item.get("amount", ""),
                        "source_table": item.get("source_table", ""),
                    }
                ]
            for idx in range(max_prices):
                price = prices[idx] if idx < len(prices) else {}
                values.extend(
                    [
                        _stringify(price.get("tax_included_price", "")),
                        _stringify(price.get("moq", "")),
                        _stringify(price.get("quantity", "")),
                        _stringify(price.get("amount", "")),
                        price.get("source_table", ""),
                    ]
                )
            values.extend(
                [
                    item.get("validation_status", ""),
                    "; ".join(item.get("errors", []) or []),
                    item.get("source", ""),
                ]
            )
            for col, value in enumerate(values, start=1):
                cell = item_sheet.cell(row=row_idx, column=col, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if col == len(values) - 2:
                    cell.fill = _status_fill(str(value))

    sheets = [ws, detail]
    if item_sheet is not None:
        sheets.append(item_sheet)

    for sheet in sheets:
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].width = 18 if col not in (6, 11, 12) else 36
        sheet.freeze_panes = "A9" if sheet is ws else "A2"

    wb.save(output_path)
    return output_path
