"""
core/excel_writer.py
====================
Excel 模板填写：将 AI 提取的字段值写入 Excel 模板。

AI 填写的内容使用红色字体，方便总务担当人工校对。
"""

import os
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import coordinate_to_tuple

from .config import BASE_DIR

# AI 填写内容使用红色（便于人工校对）
_RED_FONT = Font(color="FF0000")


def fill_excel_template(
    data: Dict[str, Any],
    cfg: Dict[str, Any],
    label_overrides: Optional[Dict[str, str]] = None,
    clear_cells: Optional[List[str]] = None,
) -> Path:
    """将提取的字段值填入 Excel 模板，返回临时输出文件路径。

    Args:
        data:            {field_id: value} 字典
        cfg:             field_mapping.json 的完整内容
        label_overrides: {cell_address: new_label} テンプレートのラベルセルを上書き
        clear_cells:     クリアするセルアドレスのリスト

    Returns:
        已填写的 Excel 文件的临时路径（调用方负责移动或删除）
    """
    raw_path = Path(cfg["excel_template_path"])
    template_path = raw_path if raw_path.is_absolute() else BASE_DIR / raw_path
    if not template_path.exists():
        raise RuntimeError(f"Excel 模板不存在: {template_path}")

    wb = load_workbook(str(template_path))

    # 选择工作表
    sheet_name = cfg.get("sheet_name")
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"Excel 中不存在工作表: {sheet_name}")
        ws = wb[sheet_name]
    else:
        ws = wb.active

    for field in cfg.get("fields", []):
        fid = field["id"]
        cell = field.get("cell", "")
        value = data.get(fid, "")
        if not value or not cell:
            continue

        # 处理合并单元格：写入合并区域的左上角单元格
        _, _ = coordinate_to_tuple(cell)  # 校验单元格地址格式
        for merged_range in ws.merged_cells.ranges:
            if cell in merged_range:
                target = ws.cell(
                    row=merged_range.min_row, column=merged_range.min_col
                )
                target.value = value
                target.font = _RED_FONT
                break
        else:
            ws[cell] = value
            ws[cell].font = _RED_FONT

    # ── ラベルセル上書き（テンプレートの固定ラベルを動的に変更）──
    for cell_addr, label_value in (label_overrides or {}).items():
        if not label_value:
            continue
        for merged_range in ws.merged_cells.ranges:
            if cell_addr in merged_range:
                target = ws.cell(
                    row=merged_range.min_row, column=merged_range.min_col
                )
                target.value = label_value
                target.font = _RED_FONT
                break
        else:
            ws[cell_addr] = label_value
            ws[cell_addr].font = _RED_FONT

    # ── セルクリア（値がない項目のラベル・値を空にする）──
    for cell_addr in (clear_cells or []):
        for merged_range in ws.merged_cells.ranges:
            if cell_addr in merged_range:
                target = ws.cell(
                    row=merged_range.min_row, column=merged_range.min_col
                )
                target.value = None
                break
        else:
            ws[cell_addr].value = None

    # 写入随机命名的临时文件，避免并发冲突
    tmp_dir = Path(tempfile.gettempdir())
    output_path = tmp_dir / f"filled_{os.getpid()}_{os.urandom(4).hex()}.xlsx"
    wb.save(str(output_path))
    return output_path
