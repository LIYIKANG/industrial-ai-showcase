"""
core/word_writer.py
===================
Word テンプレートへのデータ書き込み。

対応するプレースホルダー形式: {{field_id}}
  例: テンプレート内の「{{customer_name}}」を抽出値で置換する。

使い方:
  from core.word_writer import fill_word_template
  output_path = fill_word_template(data_dict, template_bytes, "template.docx")
"""

import re
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any, Dict


def _replace_placeholders(text: str, data: Dict[str, Any]) -> str:
    """{{field_id}} 形式のプレースホルダーを data の値で置換する。"""

    def replacer(m: re.Match) -> str:
        key = m.group(1).strip()
        return str(data.get(key, ""))

    return re.sub(r"\{\{(\w+)\}\}", replacer, text)


def fill_word_template(
    data: Dict[str, Any], template_contents: bytes, template_filename: str
) -> Path:
    """Word テンプレートにデータを書き込み、一時ファイルのパスを返す。

    Args:
        data:               フィールド ID → 値 の辞書
        template_contents:  テンプレートファイルのバイト列
        template_filename:  ファイル名（ログ用）

    Returns:
        生成された一時ファイルの Path
    """
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError(
            "python-docx がインストールされていません。"
            "pip install python-docx を実行してください。"
        )

    doc = Document(BytesIO(template_contents))

    # ── パラグラフの置換 ───────────────────────────────────────────────────────
    for para in doc.paragraphs:
        # Run を結合して全体テキストを取得し、プレースホルダーが run をまたぐ場合に対応
        full_text = "".join(run.text for run in para.runs)
        if "{{" not in full_text:
            continue
        new_text = _replace_placeholders(full_text, data)
        # 最初の run に全テキストを書き込み、残りをクリア
        if para.runs:
            para.runs[0].text = new_text
            for run in para.runs[1:]:
                run.text = ""

    # ── テーブルセルの置換 ────────────────────────────────────────────────────
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    full_text = "".join(run.text for run in para.runs)
                    if "{{" not in full_text:
                        continue
                    new_text = _replace_placeholders(full_text, data)
                    if para.runs:
                        para.runs[0].text = new_text
                        for run in para.runs[1:]:
                            run.text = ""

    # ── ヘッダー・フッターの置換 ──────────────────────────────────────────────
    for section in doc.sections:
        for hf in (section.header, section.footer):
            if hf is None:
                continue
            for para in hf.paragraphs:
                full_text = "".join(run.text for run in para.runs)
                if "{{" not in full_text:
                    continue
                new_text = _replace_placeholders(full_text, data)
                if para.runs:
                    para.runs[0].text = new_text
                    for run in para.runs[1:]:
                        run.text = ""

    # ── 一時ファイルとして保存 ────────────────────────────────────────────────
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(tmp.name)
    tmp.close()
    return Path(tmp.name)
