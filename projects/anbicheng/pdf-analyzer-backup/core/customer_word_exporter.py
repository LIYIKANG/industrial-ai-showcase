"""
core/customer_word_exporter.py
==============================
客户关键词识别结果导出为 Word (.docx)。
与 customer_exporter.py 的 Excel 导出对应，使用相同的输入参数。
"""

from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

_HEADER_BG = "E7F0EA"    # 表头底色（同 Excel）
_BOLD = True


def _set_cell(cell, text: str, *, bold: bool = False, size: int = 8):
    """设置单元格文本和基本格式。"""
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text) if text is not None else "")
    run.bold = bold
    run.font.size = Pt(size)


def _add_table(doc: Document, headers: List[str], rows: List[List[str]]) -> None:
    """添加一个带表头的表格。"""
    if not headers:
        return
    table = doc.add_table(rows=1, cols=len(headers), style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    for j, h in enumerate(headers):
        _set_cell(table.rows[0].cells[j], h, bold=True, size=8)

    # 数据行
    for row_data in rows:
        row = table.add_row()
        for j, val in enumerate(row_data):
            if j >= len(headers):
                break
            _set_cell(row.cells[j], val, bold=False, size=8)

    doc.add_paragraph()  # 表后间隔


def write_customer_result_word(
    *,
    file_id: str,
    fields: List[Dict[str, Any]],
    line_items: List[Dict[str, Any]] | None = None,
    business_result: Dict[str, Any],
    config: Dict[str, Any],
    source_filename: str,
) -> Path:
    """生成客户关键词结果 Word 文档，返回临时文件路径。"""

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)

    # ══════════════════════════════════════════════════════════
    # 标题 & 元信息
    # ══════════════════════════════════════════════════════════
    title = doc.add_paragraph()
    title.alignment = 1
    run = title.add_run("客户关键词识别结果")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1F, 0x4F, 0x3D)

    meta = doc.add_paragraph()
    meta_lines = [
        f"来源文件: {source_filename}",
        f"关键词配置: {config.get('label', config.get('id', ''))}",
        f"文件 ID: {file_id}",
        f"业务处理: {business_result.get('message', '')}",
        f"处理类型: {business_result.get('result_type', '')}",
        f"需要人工确认: {'是' if business_result.get('requires_manual_confirmation') else '否'}",
    ]
    for line in meta_lines:
        run = meta.add_run(line + "\n")
        run.font.size = Pt(9)

    doc.add_paragraph()

    # ══════════════════════════════════════════════════════════
    # 主字段表
    # ══════════════════════════════════════════════════════════
    doc.add_paragraph("主字段校对结果", style="Heading 2")
    field_headers = [
        "字段ID", "字段名称", "原始识别值", "标准化值", "校验状态", "错误/提示", "是否必填", "业务主键", "来源", "输出单元格"
    ]
    field_rows = []
    for f in fields:
        field_rows.append([
            f.get("id", ""),
            f.get("label", ""),
            str(f.get("raw_value", "")) if f.get("raw_value") is not None else "",
            str(f.get("standard_value", "")) if f.get("standard_value") is not None else "",
            f.get("validation_status", ""),
            "; ".join(f.get("errors", []) or []) or f.get("note", ""),
            "是" if f.get("required") else "否",
            "是" if f.get("business_key") else "否",
            f.get("source", ""),
            f.get("output_cell", ""),
        ])
    _add_table(doc, field_headers, field_rows)

    # ══════════════════════════════════════════════════════════
    # 业务处理
    # ══════════════════════════════════════════════════════════
    doc.add_paragraph("业务处理详情", style="Heading 2")
    biz_headers = ["项目", "值"]
    biz_rows = [
        ["主键字段", business_result.get("business_key_id", "")],
        ["主键值", business_result.get("business_key_value", "")],
        ["结果类型", business_result.get("result_type", "")],
        ["动作", business_result.get("action", "")],
        ["旧值", str(business_result.get("old_value", "")) if business_result.get("old_value") is not None else ""],
        ["新值", str(business_result.get("new_value", "")) if business_result.get("new_value") is not None else ""],
        ["匹配 MOQ", str(business_result.get("matched_moq", "")) if business_result.get("matched_moq") is not None else ""],
        ["提示", business_result.get("message", "")],
    ]
    _add_table(doc, biz_headers, biz_rows)

    # ══════════════════════════════════════════════════════════
    # SKU 明细行
    # ══════════════════════════════════════════════════════════
    if line_items:
        doc.add_paragraph("SKU 价格明细", style="Heading 2")

        # 确定价格列数（取最多的 prices 数量）
        max_prices = max(
            1,
            max(
                len(item.get("prices", []) or [])
                for item in line_items
            ),
        )

        item_headers = [
            "行号", "SKU", "客户物料代码", "品名", "规格型号/主键",
            "材质及特殊特性", "硬度", "颜色", "交货日期", "表格来源",
        ]
        for idx in range(1, max_prices + 1):
            item_headers.extend([
                f"价格{idx}", f"MOQ{idx}", f"数量{idx}", f"价税合计{idx}", f"价格来源{idx}",
            ])
        item_headers.extend(["校验状态", "错误/提示", "来源"])

        item_rows = []
        for item in line_items:
            vals = [
                item.get("line_no", ""),
                item.get("sku", ""),
                item.get("item_no", ""),
                item.get("product_name", ""),
                item.get("part_no", ""),
                item.get("material_special", ""),
                item.get("hardness", ""),
                item.get("color", ""),
                item.get("lead_time", ""),
                item.get("source_table", ""),
            ]
            prices = item.get("prices", []) or []
            if not prices:
                prices = [{
                    "tax_included_price": item.get("tax_included_price", ""),
                    "moq": item.get("moq", ""),
                    "quantity": item.get("quantity", ""),
                    "amount": item.get("amount", ""),
                    "source_table": item.get("source_table", ""),
                }]
            for idx in range(max_prices):
                price = prices[idx] if idx < len(prices) else {}
                vals.extend([
                    str(price.get("tax_included_price", "")) if price.get("tax_included_price") is not None else "",
                    str(price.get("moq", "")) if price.get("moq") is not None else "",
                    str(price.get("quantity", "")) if price.get("quantity") is not None else "",
                    str(price.get("amount", "")) if price.get("amount") is not None else "",
                    price.get("source_table", ""),
                ])
            vals.extend([
                item.get("validation_status", ""),
                "; ".join(item.get("errors", []) or []),
                item.get("source", ""),
            ])
            item_rows.append([str(v) if v is not None else "" for v in vals])

        _add_table(doc, item_headers, item_rows)

    # ══════════════════════════════════════════════════════════
    # 保存
    # ══════════════════════════════════════════════════════════
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    tmp.close()
    doc.save(tmp.name)
    return Path(tmp.name)
