"""
core/template_analyzer.py
=========================
AI によるユーザーアップロードテンプレート解析。

Excel (.xlsx) または Word (.docx) テンプレートを受け取り、
Claude AI でテンプレート構造を分析して「抽出すべきフィールド候補リスト」を返す。

返却フォーマット（各フィールド）:
  {
    "id":          "snake_case_id",
    "label":       "日本語ラベル",
    "description": "このフィールドに入れる情報の説明",
    "required":    true / false
  }
"""

import json
import re
from io import BytesIO
from typing import Any, Dict, List

from .config import CLAUDE_MODEL, claude_client

# ── Excel テキスト抽出 ─────────────────────────────────────────────────────────


def _extract_excel_text(contents: bytes) -> str:
    """Excel テンプレートの全シートからテキストを抽出する。"""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    lines = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"=== シート: {sheet_name} ===")
        for row in ws.iter_rows():
            row_texts = [str(cell.value) for cell in row if cell.value is not None]
            if row_texts:
                lines.append(" | ".join(row_texts))
    wb.close()
    return "\n".join(lines)


# ── Word テキスト抽出 ──────────────────────────────────────────────────────────


def _extract_word_text(contents: bytes) -> str:
    """Word テンプレートのパラグラフ・テーブルからテキストを抽出する。"""
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError(
            "python-docx がインストールされていません。pip install python-docx を実行してください。"
        )

    doc = Document(BytesIO(contents))
    lines = []
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                lines.append(" | ".join(row_texts))
    return "\n".join(lines)


# ── AI 解析 ───────────────────────────────────────────────────────────────────


async def analyze_template(contents: bytes, filename: str) -> List[Dict[str, Any]]:
    """テンプレートファイルを AI で解析し、抽出すべきフィールド候補リストを返す。

    Args:
        contents: テンプレートファイルのバイト列
        filename: ファイル名（拡張子で種別を判定）

    Returns:
        フィールド定義のリスト（id, label, description, required）
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in ("xlsx", "xls"):
        template_text = _extract_excel_text(contents)
        template_type = "Excel"
    elif ext in ("docx", "doc"):
        template_text = _extract_word_text(contents)
        template_type = "Word"
    else:
        raise ValueError(f"サポートされていないテンプレート形式: .{ext}（.xlsx / .docx に対応）")

    # テキストが長すぎる場合は先頭 4000 文字に絞る
    template_text_truncated = template_text[:4000]

    prompt = (
        f"以下は不動産関連書類の出力テンプレート（{template_type}形式）の内容です。\n\n"
        f"テンプレート内容:\n{template_text_truncated}\n\n"
        "このテンプレートを正しく埋めるために、PDF や画像からどのような情報を抽出する必要があるか分析してください。\n"
        "必要なフィールドのリストを以下のJSON配列形式のみで出力してください（説明文・コードブロック不要）:\n"
        '[\n'
        '  {\n'
        '    "id": "snake_case_id",\n'
        '    "label": "日本語ラベル",\n'
        '    "description": "このフィールドに入れる情報の説明（20文字以内）",\n'
        '    "required": true\n'
        '  }\n'
        ']\n'
        "フィールドは最大25個まで。JSONのみを返してください。"
    )

    resp = await claude_client.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=2048,
        system=(
            "あなたは不動産書類テンプレートの解析専門家です。"
            "指示されたJSON配列のみを出力し、余分な説明は一切含めません。"
        ),
        messages=[{"role": "user", "content": prompt}],
    )
    raw_text = resp.content[0].text.strip()

    # JSON 配列をパース（```json ... ``` ラッパーにも対応）
    cleaned = re.sub(r"^```[a-zA-Z]*\s*", "", raw_text)
    cleaned = re.sub(r"\s*```$", "", cleaned).strip()

    try:
        result = json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\[.*\]", raw_text, re.DOTALL)
        if not match:
            raise RuntimeError(
                f"テンプレート解析結果のパースに失敗しました。AI 出力:\n{raw_text[:500]}"
            )
        result = json.loads(match.group(0))

    # リスト型であることを保証
    if isinstance(result, dict):
        for v in result.values():
            if isinstance(v, list):
                result = v
                break

    if not isinstance(result, list):
        raise RuntimeError("テンプレート解析結果が配列形式ではありません。")

    # 各エントリの必須キーを補完
    sanitized = []
    for item in result:
        if not isinstance(item, dict):
            continue
        sanitized.append(
            {
                "id": str(item.get("id", f"field_{len(sanitized)}")),
                "label": str(item.get("label", item.get("id", ""))),
                "description": str(item.get("description", "")),
                "required": bool(item.get("required", True)),
            }
        )

    return sanitized
