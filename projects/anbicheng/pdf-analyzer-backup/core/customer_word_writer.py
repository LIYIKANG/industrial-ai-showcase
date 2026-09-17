"""
core/customer_word_writer.py
=============================
客户关键词识别结果导出为 Word 文档。

对标 core/customer_exporter.py 的 3 表结构，用 python-docx 从零构建。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from core.config import settings

# ── 颜色常量（与 customer_exporter.py 保持一致）──────────────────────────────────
_HEADER_BG = "E7F0EA"
_WARN_BG = "FFF4D8"
_ERROR_BG = "FFE1E1"
_OK_BG = "E7F7EF"


def _stringify(value: Any) -> str:
    """将任意值转为字符串，None 和 float 特殊处理。"""
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:g}"
    return str(value)


def _status_bg(status: str) -> str:
    """根据校验状态返回对应背景色 hex。"""
    if status == "ok":
        return _OK_BG
    if status in ("missing", "error"):
        return _ERROR_BG
    return _WARN_BG


def _set_cell_shading(cell, color_hex: str) -> None:
    """设置 Word 表格单元格背景色。"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    tcPr.append(shading)


def _write_cell(cell, text: str, *, bold: bool = False, shading: str | None = None,
                size: int = 9) -> None:
    """写入单元格文本并应用格式。"""
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if shading:
        _set_cell_shading(cell, shading)


def _add_info_para(doc: Document, label: str, value: str) -> None:
    """添加「标签：值」格式的元信息段落。"""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(2)
    run_label = para.add_run(f"{label}：")
    run_label.bold = True
    run_label.font.size = Pt(10)
    run_value = para.add_run(value)
    run_value.font.size = Pt(10)


def _add_section_heading(doc: Document, title: str) -> None:
    """添加段落标题（一、校对结果 等）。"""
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(title)
    run.bold = True
    run.font.size = Pt(12)


def _build_verification_table(doc: Document, fields: List[Dict[str, Any]]) -> None:
    """构建「校对结果」表格（10 列）。"""
    headers = [
        "字段ID", "字段名称", "原始识别值", "标准化值", "校验状态",
        "错误/提示", "是否必填", "业务主键", "来源", "输出单元格",
    ]
    num_cols = len(headers)
    num_rows = 1 + len(fields)  # 表头 + 数据行
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = "Table Grid"

    # 表头行
    for col, header in enumerate(headers):
        _write_cell(table.cell(0, col), header, bold=True, shading=_HEADER_BG)

    # 数据行
    for row_idx, field in enumerate(fields, start=1):
        values = [
            _stringify(field.get("id", "")),
            _stringify(field.get("label", "")),
            _stringify(field.get("raw_value", "")),
            _stringify(field.get("standard_value", "")),
            _stringify(field.get("validation_status", "")),
            "; ".join(field.get("errors", []) or []) or _stringify(field.get("note", "")),
            "是" if field.get("required") else "否",
            "是" if field.get("business_key") else "否",
            _stringify(field.get("source", "")),
            _stringify(field.get("output_cell", "")),
        ]
        for col, value in enumerate(values):
            shading = _status_bg(str(field.get("validation_status", ""))) if col == 4 else None
            _write_cell(table.cell(row_idx, col), value, shading=shading)


def _build_business_table(doc: Document, business_result: Dict[str, Any]) -> None:
    """构建「业务处理」表格（2 列 x 8 行）。"""
    rows_data = [
        ("主键字段", _stringify(business_result.get("business_key_id", ""))),
        ("主键值", _stringify(business_result.get("business_key_value", ""))),
        ("结果类型", _stringify(business_result.get("result_type", ""))),
        ("动作", _stringify(business_result.get("action", ""))),
        ("旧值", _stringify(business_result.get("old_value", ""))),
        ("新值", _stringify(business_result.get("new_value", ""))),
        ("匹配 MOQ", _stringify(business_result.get("matched_moq", ""))),
        ("提示", _stringify(business_result.get("message", ""))),
    ]
    table = doc.add_table(rows=len(rows_data) + 1, cols=2)
    table.style = "Table Grid"

    # 表头
    for col, header in enumerate(["项目", "值"]):
        _write_cell(table.cell(0, col), header, bold=True, shading=_HEADER_BG)

    # 数据行
    for row_idx, (label, value) in enumerate(rows_data, start=1):
        _write_cell(table.cell(row_idx, 0), label, bold=True)
        _write_cell(table.cell(row_idx, 1), value)


def _build_line_items_table(doc: Document, line_items: List[Dict[str, Any]]) -> None:
    """构建「明细行」表格（动态列数）。"""
    if not line_items:
        return

    # 计算最大价格组数（与 customer_exporter.py 逻辑一致）
    raw_max = max(
        (len(item.get("prices", []) or []) for item in line_items),
        default=1,
    )
    max_prices = max(1, raw_max)

    # 固定列
    fixed_headers = [
        "行号", "SKU", "客户物料代码", "品名", "规格型号/主键",
        "材质及特殊特性", "硬度", "颜色", "交货日期", "表格来源",
    ]
    # 动态价格组
    price_headers = []
    for idx in range(1, max_prices + 1):
        price_headers.extend([
            f"价格{idx}", f"MOQ{idx}", f"数量{idx}", f"价税合计{idx}", f"价格来源{idx}",
        ])
    trailing_headers = ["校验状态", "错误/提示", "来源"]
    all_headers = fixed_headers + price_headers + trailing_headers

    num_cols = len(all_headers)
    num_rows = 1 + len(line_items)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = "Table Grid"

    # 表头行（小字号适配横版多列）
    for col, header in enumerate(all_headers):
        _write_cell(table.cell(0, col), header, bold=True, shading=_HEADER_BG, size=8)

    # 数据行
    for row_idx, item in enumerate(line_items, start=1):
        # 校验状态列的位置（用于颜色编码）
        status_col = 10 + 5 * max_prices  # 固定10列 + N组价格列

        fixed_values = [
            _stringify(item.get("line_no", "")),
            _stringify(item.get("sku", "")),
            _stringify(item.get("item_no", "")),
            _stringify(item.get("product_name", "")),
            _stringify(item.get("part_no", "")),
            _stringify(item.get("material_special", "")),
            _stringify(item.get("hardness", "")),
            _stringify(item.get("color", "")),
            _stringify(item.get("lead_time", "")),
            _stringify(item.get("source_table", "")),
        ]

        prices = item.get("prices", []) or []
        if not prices:
            prices = [{
                "tax_included_price": item.get("tax_included_price", ""),
                "moq": item.get("moq", ""),
                "quantity": item.get("quantity", ""),
                "amount": item.get("amount", ""),
                "source_table": item.get("source_table", ""),
            }]

        price_values = []
        for idx in range(max_prices):
            price = prices[idx] if idx < len(prices) else {}
            price_values.extend([
                _stringify(price.get("tax_included_price", "")),
                _stringify(price.get("moq", "")),
                _stringify(price.get("quantity", "")),
                _stringify(price.get("amount", "")),
                _stringify(price.get("source_table", "")),
            ])

        trailing_values = [
            _stringify(item.get("validation_status", "")),
            "; ".join(item.get("errors", []) or []),
            _stringify(item.get("source", "")),
        ]

        all_values = fixed_values + price_values + trailing_values
        validation_status = _stringify(item.get("validation_status", ""))

        for col, value in enumerate(all_values):
            shading = _status_bg(validation_status) if col == status_col else None
            _write_cell(table.cell(row_idx, col), value, shading=shading, size=8)


def write_customer_result_word(
    *,
    file_id: str,
    fields: List[Dict[str, Any]],
    line_items: List[Dict[str, Any]] | None = None,
    business_result: Dict[str, Any],
    config: Dict[str, Any],
    source_filename: str,
) -> Path:
    """生成客户关键词结果 Word 文档，返回文件路径。

    文档结构：
      - 标题：客户关键词识别结果
      - 基本信息（来源文件、关键词配置、业务处理等）
      - 表格1：校对结果（字段级表格）
      - 表格2：业务处理（键值表格）
      - 表格3：明细行（仅在 line_items 非空时出现，动态列数）
    """
    settings.CUSTOMER_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = settings.CUSTOMER_EXPORT_DIR / f"{file_id}.docx"

    doc = Document()

    # ── 标题 ──────────────────────────────────────────────────────────────────
    title_para = doc.add_paragraph()
    title_para.alignment = 1  # 居中
    title_run = title_para.add_run("客户关键词识别结果")
    title_run.bold = True
    title_run.font.size = Pt(16)

    # ── 元信息 ────────────────────────────────────────────────────────────────
    _add_info_para(doc, "来源文件", source_filename)
    _add_info_para(doc, "关键词配置", _stringify(config.get("label", config.get("id", ""))))
    _add_info_para(doc, "业务处理", _stringify(business_result.get("message", "")))
    _add_info_para(doc, "处理类型", _stringify(business_result.get("result_type", "")))
    _add_info_para(
        doc,
        "是否需要人工确认",
        "是" if business_result.get("requires_manual_confirmation") else "否",
    )

    # ── 表格1：校对结果 ───────────────────────────────────────────────────────
    _add_section_heading(doc, "一、校对结果")
    _build_verification_table(doc, fields)

    # ── 表格2：业务处理 ───────────────────────────────────────────────────────
    _add_section_heading(doc, "二、业务处理")
    _build_business_table(doc, business_result)

    # ── 表格3：明细行（可选，横版页面）──────────────────────────────────────────
    if line_items:
        # 切换到横版（landscape）节，给多列表格更多空间
        new_section = doc.add_section()
        new_section.orientation = WD_ORIENT.LANDSCAPE
        new_section.page_width = Cm(29.7)
        new_section.page_height = Cm(21.0)
        # 缩窄页边距以最大化表格宽度
        new_section.top_margin = Cm(1.2)
        new_section.bottom_margin = Cm(1.2)
        new_section.left_margin = Cm(1.5)
        new_section.right_margin = Cm(1.5)

        _add_section_heading(doc, "三、明细行")
        _build_line_items_table(doc, line_items)

    doc.save(str(output_path))
    return output_path
