"""
scripts/generate_customer_word_template.py
==========================================
根据 config/customer_keywords.json 自动生成客户关键词 Word 模板 (.docx)。
模板包含主字段表和 SKU 明细表，使用 {{field_id}} 占位符。
"""

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE_DIR = Path(__file__).resolve().parent.parent


def generate(config_path: str, output_path: str) -> Path:
    config_file = BASE_DIR / config_path
    with open(config_file, "r", encoding="utf-8") as f:
        cfg = json.load(f)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)

    # ── 标题 ──
    title = doc.add_paragraph()
    title.alignment = 1
    run = title.add_run(cfg.get("label", "客户关键词识别"))
    run.bold = True
    run.font.size = Pt(16)

    doc.add_paragraph()

    # ── Sheet 1：主字段表 ──
    fields = cfg.get("fields", [])
    doc.add_paragraph("主字段", style="Heading 2")
    table = doc.add_table(rows=1, cols=2, style="Table Grid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "字段"
    hdr[1].text = "值（{{field_id}}）"
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)

    for f in fields:
        row = table.add_row()
        row.cells[0].text = f.get("label", f["id"])
        row.cells[1].text = "{{" + f["id"] + "}}"

    doc.add_paragraph()

    # ── Sheet 2：SKU 明细表 ──
    columns = cfg.get("line_items", {}).get("columns", [])
    if columns:
        doc.add_paragraph("SKU 价格明细", style="Heading 2")
        item_table = doc.add_table(rows=1, cols=len(columns), style="Table Grid")
        item_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr2 = item_table.rows[0].cells
        for j, col in enumerate(columns):
            hdr2[j].text = col.get("label", col["id"])
            for r in hdr2[j].paragraphs:
                for run in r.runs:
                    run.bold = True
                    run.font.size = Pt(8)

        # 5 行示例占位符
        for i in range(5):
            row = item_table.add_row()
            for j, col in enumerate(columns):
                row.cells[j].text = "{{line_" + col["id"] + "_" + str(i + 1) + "}}"

    output = BASE_DIR / output_path
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    print(f"Word 模板已生成: {output}")
    return output


if __name__ == "__main__":
    generate("config/customer_keywords.json", "data/templates/客户关键词模板.docx")
