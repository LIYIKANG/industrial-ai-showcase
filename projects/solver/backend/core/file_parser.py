from __future__ import annotations
import io
import pandas as pd
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

def _decode(data):
    for enc in ("utf-8-sig","utf-8","gbk","gb18030"):
        try:
            return data.decode(enc)
        except Exception:
            pass
    return data.decode("utf-8", errors="ignore")

def parse_file_bytes(filename: str, data: bytes) -> str:
    name=(filename or "").lower()
    if not data: raise ValueError("文件为空。")
    if name.endswith((".txt",".json")): return _decode(data)
    if name.endswith(".pdf"):
        reader=PdfReader(io.BytesIO(data))
        return "\n\n".join([f"【Page {i+1}】\n{p.extract_text() or ''}" for i,p in enumerate(reader.pages)])
    if name.endswith(".pptx"):
        prs=Presentation(io.BytesIO(data))
        out=[]
        for i,sl in enumerate(prs.slides,1):
            out.append(f"【Slide {i}】")
            for sh in sl.shapes:
                if hasattr(sh,"text") and sh.text.strip(): out.append(sh.text.strip())
                if getattr(sh, "has_table", False):
                    for row in sh.table.rows:
                        values=[cell.text.strip() for cell in row.cells]
                        if any(values): out.append(" | ".join(values))
        return "\n".join(out)
    if name.endswith(".docx"):
        doc=Document(io.BytesIO(data)); out=[p.text for p in doc.paragraphs if p.text.strip()]
        for tb in doc.tables:
            for row in tb.rows:
                vals=[c.text.strip() for c in row.cells]
                if any(vals): out.append(" | ".join(vals))
        return "\n".join(out)
    if name.endswith((".xlsx",".xls")):
        xls=pd.ExcelFile(io.BytesIO(data)); chunks=[]
        for s in xls.sheet_names:
            df=pd.read_excel(xls,sheet_name=s)
            chunks += [f"【Sheet: {s}】", df.head(1000).to_csv(index=False)]
        return "\n".join(chunks)
    if name.endswith(".csv"):
        for enc in ("utf-8-sig","utf-8","gbk","gb18030"):
            try: return pd.read_csv(io.BytesIO(data),encoding=enc).head(2000).to_csv(index=False)
            except Exception: pass
        return _decode(data)
    raise ValueError("支持 PDF/PPTX/DOCX/XLSX/CSV/TXT/JSON。")
